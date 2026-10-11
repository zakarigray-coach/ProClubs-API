import io
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock
from datetime import datetime
import openpyxl
from gaffergpt.excel_intake import read_excel
from gaffergpt.bot import GafferBot
from gaffergpt.config import Config

def workbook_bytes():
    book=openpyxl.Workbook();s=book.active;s.title='Matches'
    s.append(['Date','Time (UTC)','Match ID','Match type','Goals for','Goals against','Opponent'])
    s.append([datetime(2026,10,7,2,41),'02:41','123','league',2,2,'Chelsea'])
    p=book.create_sheet('Players');p.append(['Player','Goals','Assists','Pass %']);p.append(['Zak',2,3,0.75])
    p=book.create_sheet('Player matches');p.append(['Player','Match ID','Rating']);p.append(['Zak','123',7.5])
    out=io.BytesIO();book.save(out);return out.getvalue()

class ExcelReadTests(unittest.TestCase):
    def test_all_sheets_original_values_and_explicit_utc(self):
        result=json.loads(read_excel(workbook_bytes()))
        self.assertEqual(len(result['worksheets']),3)
        self.assertIn('+00:00',result['worksheets'][0]['rows'][1][0])
        self.assertEqual(result['worksheets'][1]['rows'][1],['Zak',2,3,0.75])
        self.assertIn('not official',result['date_note'])
    def test_invalid_oversized_and_empty_workbook(self):
        for data in [b'not excel',b'x'*(8*1024*1024+1)]:
            with self.assertRaises(ValueError): read_excel(data)
        book=openpyxl.Workbook();out=io.BytesIO();book.save(out)
        with self.assertRaises(ValueError):read_excel(out.getvalue())
    def test_export_without_dimension_metadata(self):
        import zipfile,re
        original=io.BytesIO(workbook_bytes());out=io.BytesIO()
        with zipfile.ZipFile(original) as src,zipfile.ZipFile(out,'w') as dst:
            for item in src.infolist():
                data=src.read(item.filename)
                if item.filename.startswith('xl/worksheets/'):
                    data=re.sub(rb'<dimension[^>]*/>',b'',data)
                dst.writestr(item,data)
        result=json.loads(read_excel(out.getvalue()))
        self.assertEqual(len(result['worksheets'][0]['rows']),2)

    def test_formula_never_executes(self):
        book=openpyxl.Workbook();book.active.append(['Player','Value']);book.active.append(['Zak','=1+1'])
        out=io.BytesIO();book.save(out)
        self.assertIsNone(json.loads(read_excel(out.getvalue()))['worksheets'][0]['rows'][1][1])

class ExcelDMTests(unittest.IsolatedAsyncioTestCase):
    async def test_upload_then_classification_reply_reuses_excel(self):
        with tempfile.TemporaryDirectory() as root:
            bot=GafferBot(Config('test',1,2,database_path=str(Path(root)/'pep.db')))
            bot.recap_reader=SimpleNamespace(extract=AsyncMock(return_value={'lines':[],'questions':['Official or practice?']}))
            class Typing:
                async def __aenter__(self):pass
                async def __aexit__(self,*args):pass
            data=workbook_bytes()
            msg=SimpleNamespace(id=20,guild=None,author=SimpleNamespace(id=1,bot=False),content='',attachments=[SimpleNamespace(filename='pct.xlsx',size=len(data),read=AsyncMock(return_value=data))],reply=AsyncMock(),channel=SimpleNamespace(typing=lambda:Typing()))
            await bot.on_message(msg)
            self.assertIn('no need to upload',msg.reply.call_args.args[0])
            self.assertIn('Player matches',bot.recap_reader.extract.call_args.args[0])
            self.assertFalse(bot.db.rows('SELECT * FROM matches'))
            msg.id=21;msg.content='All these matches were practices.';msg.attachments=[]
            bot.recap_reader.extract.return_value={'lines':['2026-10-06 | Practice | Chelsea | 2-2'],'questions':[]}
            await bot.on_message(msg)
            self.assertIn('All these matches were practices',bot.recap_reader.extract.call_args.args[0])
            self.assertIn('Zak',bot.recap_reader.extract.call_args.args[0])
            self.assertTrue(msg.reply.call_args.kwargs['view'].is_persistent())
            self.assertFalse(bot.db.rows('SELECT * FROM matches'))
            await bot.close()
