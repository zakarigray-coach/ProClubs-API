"""Read XLSX values without executing formulas, macros, or external links."""
import io
import json
import zipfile
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
import openpyxl

MAX_BYTES=8*1024*1024

def read_excel(data):
    if len(data)>MAX_BYTES: raise ValueError('Excel file must be 8 MB or smaller.')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist())>40*1024*1024 or len(archive.infolist())>200:
                raise ValueError('This workbook is too large to read safely.')
            if any('vbaProject' in name for name in archive.namelist()):
                raise ValueError('Please send a normal .xlsx file without macros.')
        book=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False)
        sheets=[];cells=0
        try:
            for sheet in book:
                if (sheet.max_row or 0)>1500 or (sheet.max_column or 0)>80:
                    raise ValueError('This workbook has too many rows or columns. Export a shorter reporting period.')
                rows=[]
                for row_number,row in enumerate(sheet.iter_rows(values_only=True),1):
                    if row_number>1500 or len(row)>80: raise ValueError('This workbook has too many rows or columns.')
                    cells+=len(row)
                    if cells>40000: raise ValueError('This workbook has too many cells.')
                    converted=[]
                    for value in row:
                        if isinstance(value,datetime):
                            value=value.replace(tzinfo=timezone.utc).isoformat() if value.tzinfo is None else value.isoformat()
                        if isinstance(value,str) and len(value)>2000: raise ValueError('A workbook cell is too long.')
                        converted.append(value)
                    if any(v is not None for v in converted): rows.append(converted)
                sheets.append({'sheet':sheet.title,'rows':rows})
        finally: book.close()
    except ValueError: raise
    except Exception as error:
        raise ValueError('Could not read this Excel file. Send an unencrypted .xlsx export.') from error
    if not any(s['rows'] for s in sheets): raise ValueError('The workbook has no readable data.')
    result=json.dumps({'source':'Owner-uploaded Excel export','date_note':'PCT+ Date and Time (UTC) are UTC. Use America/New_York calendar dates for reporting. EA league/friendly is not official MPL classification. Blank cells mean unavailable, not zero. Weekly player totals apply to the exported period only.','worksheets':sheets},ensure_ascii=False,default=str,separators=(',',':'))
    if len(result)>180000: raise ValueError('Export a shorter reporting period; this workbook is too large.')
    return result
